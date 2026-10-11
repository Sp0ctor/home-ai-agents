#!/usr/bin/env python3
"""Turn Athena's Markdown draft into finished files the owner can submit.

  make_doc.py DRAFT.md [--apa] [--title "Title"] [--out FOLDER]
        -> DRAFT.docx + DRAFT.pdf  (Word + PDF, same name, next to the draft or in FOLDER)
  make_doc.py slides OUTLINE.md [--out FOLDER]
        -> OUTLINE.pptx + OUTLINE.pdf  (one slide per "## " heading; bullets under it; "Notes:" lines = speaker notes)
  make_doc.py pdf FILE.docx|FILE.pptx
        -> FILE.pdf  (convert any Word/PowerPoint file to PDF)

Markdown it understands: # Title, ## Heading, ### Subheading, - bullets, 1. numbered, **bold**, *italic*,
| tables |, > quotes, blank line = new paragraph. "References" heading -> hanging-indent APA list.
--apa  : Times New Roman 12, double-spaced, 1" margins, page numbers, APA title page from the first lines:
         # Title / then lines "Author: ...", "Course: ...", "Instructor: ...", "Date: ..." right after the title.

Checks its own work: prints RESULT: OK with sizes and page count only if both files exist and reopen.
Run with /Users/you/.hermes/hermes-agent/venv/bin/python.
"""
import re, subprocess, sys, tempfile, zipfile, shutil
from pathlib import Path

HERE = Path(__file__).resolve().parent
SOFFICE = Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"
if not SOFFICE.exists():
    SOFFICE = Path(shutil.which("soffice") or "soffice")


def to_pdf(path, outdir):
    r = subprocess.run([str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
                       capture_output=True, text=True, timeout=240)
    pdf = Path(outdir) / (Path(path).stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("PDF conversion failed: " + (r.stderr or r.stdout)[-300:])
    return pdf


def pdf_pages(pdf):
    try:
        import pypdfium2 as pdfium
        return len(pdfium.PdfDocument(str(pdf)))
    except Exception:
        from pypdf import PdfReader
        return len(PdfReader(str(pdf)).pages)


# ---------------------------------------------------------------- inline formatting
def add_runs(par, text, size=None, font=None):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Courier New"
        else:
            r = par.add_run(part)
        if size:
            r.font.size = size
        if font and not part.startswith("`"):
            r.font.name = font


# ---------------------------------------------------------------- Word
def build_docx(md_path, out_dir, apa=False, title=None):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches

    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    d = docx.Document()
    font = "Times New Roman" if apa else "Calibri"
    st = d.styles["Normal"]
    st.font.name, st.font.size = font, Pt(12 if apa else 11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    pf = st.paragraph_format
    pf.space_after = Pt(0 if apa else 6)
    if apa:
        pf.line_spacing = 2.0
    for s in d.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    for lvl in (1, 2, 3):
        hs = d.styles[f"Heading {lvl}"]
        hs.font.name, hs.font.color.rgb = font, None
        rf = hs.element.get_or_add_rPr().get_or_add_rFonts()
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(att), None)
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), font)
        hs.font.size = Pt(12 if apa else (16, 13, 12)[lvl - 1])
        hs.font.bold = True
        hs.font.italic = apa and lvl == 3
        hs.paragraph_format.space_before = Pt(0 if apa else 12)
        hs.paragraph_format.space_after = Pt(0 if apa else 4)
        if apa:
            hs.paragraph_format.line_spacing = 2.0
            if lvl == 1:
                hs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if apa:  # page number top right
        hp = d.sections[0].header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)

    i = 0
    doc_title = title
    meta = {}
    if lines and lines[0].startswith("# "):
        doc_title = doc_title or lines[0][2:].strip()
        i = 1
        while i < len(lines) and (not lines[i].strip() or re.match(r"^(Author|Name|Course|Instructor|Professor|Date|Institution|School)\s*:", lines[i], re.I)):
            m = re.match(r"^(\w+)\s*:\s*(.+)$", lines[i].strip())
            if m:
                meta[m.group(1).lower()] = m.group(2)
            i += 1
    if apa and doc_title:
        for _ in range(3):
            d.add_paragraph()
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(doc_title); r.bold = True
        d.add_paragraph()
        for k in ("author", "name", "institution", "school", "course", "instructor", "professor", "date"):
            if k in meta:
                p = d.add_paragraph(meta[k]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(doc_title); r.bold = True
    elif doc_title:
        d.add_heading(doc_title, 0).runs[0].font.name = font
        if meta:
            d.add_paragraph(" | ".join(meta.values())).runs[0].italic = True

    in_refs = False
    para = []

    def flush():
        if para:
            p = d.add_paragraph()
            add_runs(p, " ".join(x.strip() for x in para))
            if apa and not in_refs:
                p.paragraph_format.first_line_indent = Inches(0.5)
            if in_refs:
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
            para.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            flush()
        elif s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            n = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=n)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(n):
                    cell = t.cell(ri, ci)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                    add_runs(cell.paragraphs[0], row[ci] if ci < len(row) else "", Pt(10 if not apa else 11))
                    if ri == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            d.add_paragraph()
            continue
        elif re.match(r"^#{1,3} ", s):
            flush()
            lvl = len(s.split(" ")[0])
            text = s[lvl + 1:].strip()
            in_refs = bool(re.match(r"^(references|works cited|sources|bibliography)$", text, re.I))
            h = d.add_heading(level=max(1, lvl - 1) if apa else min(lvl, 3))
            add_runs(h, text)
        elif re.match(r"^[-*+] ", s):
            flush()
            p = d.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
            if apa:
                p.paragraph_format.line_spacing = 2.0
        elif re.match(r"^\d+[.)] ", s):
            flush()
            p = d.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+[.)] ", "", s))
        elif s.startswith(">"):
            flush()
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            add_runs(p, s.lstrip("> "))
        elif s in ("---", "***", "<pagebreak>"):
            flush()
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            para.append(s)
            if in_refs:
                flush()
        i += 1
    flush()
    out = Path(out_dir) / (Path(md_path).stem + ".docx")
    d.save(out)
    return out


# ---------------------------------------------------------------- slides
def pptx_lib():
    vend = HERE / "vendor"
    if not (vend / "pptx").is_dir():
        zipfile.ZipFile(HERE / "pptx_lib.zip").extractall(vend)
    sys.path.insert(0, str(vend))
    import pptx
    return pptx


def build_pptx(md_path, out_dir):
    pptx = pptx_lib()
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    NAVY, ACCENT, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2E, 0x75, 0xB6), RGBColor(0x40, 0x40, 0x40)

    text = Path(md_path).read_text(encoding="utf-8")
    title, subtitle, slides, cur = None, [], [], None
    for ln in text.splitlines():
        s = ln.strip()
        if s.startswith("# ") and title is None:
            title = s[2:].strip()
        elif s.startswith("## "):
            cur = {"title": s[3:].strip(), "bullets": [], "notes": []}
            slides.append(cur)
        elif cur is None and s and title:
            subtitle.append(s)
        elif cur is not None and re.match(r"^(notes?|speaker notes?)\s*:", s, re.I):
            cur["notes"].append(s.split(":", 1)[1].strip())
        elif cur is not None and cur["notes"] and s and not re.match(r"^[-*+] |^\d+[.)] ", s):
            cur["notes"].append(s)
        elif cur is not None and s:
            lvl = 1 if re.match(r"^\s{2,}[-*+] ", ln) else 0
            cur["bullets"].append((lvl, re.sub(r"^([-*+]|\d+[.)])\s+", "", s)))

    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def bar(slide):
        shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
        shp.fill.solid(); shp.fill.fore_color.rgb = ACCENT; shp.line.fill.background()

    s0 = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s0.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    tb = s0.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6)).text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]; p.text = title or Path(md_path).stem
    p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.color.rgb = Pt(40), True, RGBColor(255, 255, 255)
    for line in subtitle[:4]:
        q = tb.add_paragraph(); q.text = re.sub(r"\*\*?", "", line)
        q.runs[0].font.size, q.runs[0].font.color.rgb = Pt(20), RGBColor(0xDD, 0xE6, 0xF0)

    for sd in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bar(sl)
        t = sl.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(1)).text_frame
        t.word_wrap = True
        t.paragraphs[0].text = sd["title"]
        r = t.paragraphs[0].runs[0]; r.font.size, r.font.bold, r.font.color.rgb = Pt(32), True, NAVY
        n = len(sd["bullets"])
        size = 24 if n <= 4 else 20 if n <= 6 else 17
        body = sl.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.6), Inches(5.4)).text_frame
        body.word_wrap = True
        for k, (lvl, b) in enumerate(sd["bullets"]):
            para = body.paragraphs[0] if k == 0 else body.add_paragraph()
            para.level = lvl
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", b)
            bold = b.startswith("**")
            run = para.add_run(); run.text = ("• " if lvl == 0 else "– ") + clean
            run.font.size, run.font.color.rgb, run.font.bold = Pt(size - 4 * lvl), GREY, bold
            para.space_after = Pt(10)
        if sd["notes"]:
            sl.notes_slide.notes_text_frame.text = " ".join(sd["notes"])
    out = Path(out_dir) / (Path(md_path).stem + ".pptx")
    prs.save(out)
    return out, len(slides) + 1


# ---------------------------------------------------------------- main
def main(a):
    if not a:
        print(__doc__); return
    out_dir = None
    if "--out" in a:
        k = a.index("--out"); out_dir = Path(a[k + 1]).expanduser(); del a[k:k + 2]
    title = None
    if "--title" in a:
        k = a.index("--title"); title = a[k + 1]; del a[k:k + 2]
    apa = "--apa" in a
    a = [x for x in a if x != "--apa"]
    mode = a[0] if a[0] in ("slides", "pdf") else "doc"
    src = Path(a[1] if mode != "doc" else a[0]).expanduser().resolve()
    if not src.exists():
        sys.exit(f"RESULT: FAIL - no such file {src}")
    out_dir = out_dir or src.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    if mode == "pdf":
        made = [to_pdf(src, out_dir)]
    elif mode == "slides":
        pp, n = build_pptx(src, out_dir)
        made = [pp, to_pdf(pp, out_dir)]
    else:
        dx = build_docx(src, out_dir, apa, title)
        made = [dx, to_pdf(dx, out_dir)]
    for f in made:  # verify: exists, non-empty, reopens
        if not f.exists() or f.stat().st_size == 0:
            sys.exit(f"RESULT: FAIL - {f.name} was not created")
        if f.suffix == ".docx":
            import docx; docx.Document(str(f))
        elif f.suffix == ".pptx":
            pptx_lib().Presentation(str(f))
    pdf = next(f for f in made if f.suffix == ".pdf")
    print(f"RESULT: OK - {pdf_pages(pdf)} page(s)")
    for f in made:
        print(f"  {f}  ({f.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except SystemExit:
        raise
    except Exception as ex:
        print(f"RESULT: FAIL - {type(ex).__name__}: {ex}")
        sys.exit(1)
