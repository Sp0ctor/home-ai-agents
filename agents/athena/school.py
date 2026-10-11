#!/usr/bin/env python3
"""Athena's course library. One place for every school file, readable and searchable.

  school.py sync                      file every new upload (Telegram files, zips, old folders) into the library
  school.py read FILE [PART]          print the text of any PDF / Word / PowerPoint / Excel / CSV / text / zip
                                      (long files come in parts: PART = 2, 3, ...)
  school.py search "words" [COURSE]   find where the course files talk about something (file + page + snippet)
  school.py list [COURSE]             what's in the library
  school.py file FILE COURSE [FOLDER] put a file in a course (or move a wrongly-filed one)
  school.py index                     rebuild the 00 Index.md pages (sync does this for you)

Library: ~/Documents/Athena/<COURSE>/{Syllabus, Course Materials, My Work, Notes}
Courses: set them in COURSES below (anything unclear goes to _Unsorted).
Run with /Users/you/.hermes/hermes-agent/venv/bin/python. Never deletes anything.
"""
import warnings
warnings.filterwarnings("ignore")
import datetime as dt, hashlib, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

HOME = Path.home()
PROFILE = HOME / ".hermes/profiles/athena"
LIB = HOME / "Documents/Athena"
TEXT = LIB / ".text"
REG = LIB / ".library.json"
SOURCES = [PROFILE / "cache/documents", PROFILE / "cache/images", PROFILE / "audio_cache", PROFILE / "Athena/vault", PROFILE / "School/courses"]
SOFFICE = HOME / "Applications/LibreOffice.app/Contents/MacOS/soffice"
FOLDERS = ["Syllabus", "Course Materials", "My Work", "Notes"]
PART_CHARS = 12000
DOC_EXT = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".csv", ".txt", ".md", ".rtf", ".odt", ".sql", ".json",
           ".html", ".htm"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".heic", ".webp", ".gif"}          # photos/screenshots sent on Telegram
AUDIO_EXT = {".ogg", ".oga", ".m4a", ".mp3", ".wav", ".mp4", ".mov"}  # voice notes / recordings

COURSES = {   # one entry per course: folder name -> full name, a regex for its code, and words that only show up in it
    "CS-101": {"name": "Intro to Computer Science",
               "code": r"\bCS[\s._-]?101\b",
               "hints": r"python|algorithm|big-o|recursion"},
    "BUS-210": {"name": "IT Management",
                "code": r"\bBUS[\s._-]?210\b",
                "hints": r"\bCIO\b|governance|IT strategy|case study"},
    "DB-304": {"name": "SQL and NoSQL Databases",
               "code": r"\bDB[\s._-]?304\b",
               "hints": r"\bSELECT\b.*\bFROM\b|NoSQL|normaliz|ERD"},
}


# ---------------------------------------------------------------- reading any file
def _xml_text(data, tag="t"):
    out = []
    for el in ET.fromstring(data).iter():
        if el.tag.endswith("}" + tag) and el.text:
            out.append(el.text)
    return out


def _pdf(path):
    pages = []
    for reader in ("pdfium", "pypdf", "plumber"):
        try:
            if reader == "pdfium":
                import pypdfium2 as pdfium
                doc = pdfium.PdfDocument(str(path))
                pages = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
            elif reader == "pypdf":
                from pypdf import PdfReader
                pages = [(p.extract_text() or "") for p in PdfReader(str(path)).pages]
            else:
                import pdfplumber
                with pdfplumber.open(str(path)) as pdf:
                    pages = [(p.extract_text() or "") for p in pdf.pages]
        except Exception:
            pages = []
        if "".join(pages).strip():
            break
    if not "".join(pages).strip():
        return "(This PDF has no text layer - it is probably a scan or images. Ask the owner for a text version.)"
    return "\n".join(f"=== page {i} ===\n{t.replace(chr(13), '').strip()}" for i, t in enumerate(pages, 1))


def _docx(path):
    import docx
    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                if c.text.strip() and (not cells or cells[-1] != c.text.strip()):
                    cells.append(c.text.strip())
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _pptx(path):
    z = zipfile.ZipFile(path)
    slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                    key=lambda n: int(re.search(r"\d+", n.rsplit("/", 1)[1]).group()))
    out = []
    for n in slides:
        k = re.search(r"(\d+)\.xml", n).group(1)
        out.append(f"=== slide {k} ===\n" + "\n".join(_xml_text(z.read(n))))
        note = f"ppt/notesSlides/notesSlide{k}.xml"
        if note in z.namelist():
            txt = [t for t in _xml_text(z.read(note)) if not t.strip().isdigit()]
            if txt:
                out.append("Speaker notes: " + " ".join(txt))
    return "\n".join(out)


def _xlsx(path):
    z = zipfile.ZipFile(path)
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root:
            shared.append("".join(t.text or "" for t in si.iter() if t.tag.endswith("}t")))
    names = []
    if "xl/workbook.xml" in z.namelist():
        names = [s.get("name") for s in ET.fromstring(z.read("xl/workbook.xml")).iter() if s.tag.endswith("}sheet")]
    out = []
    sheets = sorted((n for n in z.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)),
                    key=lambda n: int(re.search(r"(\d+)\.xml", n).group(1)))
    for idx, n in enumerate(sheets):
        out.append(f"=== sheet {names[idx] if idx < len(names) else idx + 1} ===")
        for row in ET.fromstring(z.read(n)).iter():
            if not row.tag.endswith("}row"):
                continue
            vals = []
            for c in row:
                v = next((x.text for x in c if x.tag.endswith("}v")), None)
                if v is None:
                    v = "".join(t.text or "" for t in c.iter() if t.tag.endswith("}t"))
                elif c.get("t") == "s":
                    v = shared[int(v)] if int(v) < len(shared) else v
                vals.append(v or "")
            if any(vals):
                out.append(" | ".join(vals))
    return "\n".join(out)


def _convert(path, fmt):
    """Old Office formats -> new, with LibreOffice. Returns the converted file or None."""
    if not SOFFICE.exists():
        return None
    tmp = Path(tempfile.mkdtemp())
    subprocess.run([str(SOFFICE), "--headless", "--convert-to", fmt, "--outdir", str(tmp), str(path)],
                   capture_output=True, timeout=180)
    got = list(tmp.glob(f"*.{fmt}"))
    return got[0] if got else None


def _html(path):
    """Course readings saved from the portal as .html -> plain text (headings and paragraphs kept)."""
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.out, self.skip = [], 0

        def handle_starttag(self, tag, a):
            if tag in ("script", "style", "noscript"):
                self.skip += 1
            elif tag in ("p", "br", "li", "h1", "h2", "h3", "h4", "tr", "div"):
                self.out.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "noscript") and self.skip:
                self.skip -= 1

        def handle_data(self, d):
            if not self.skip:
                self.out.append(d)

    p = P()
    p.feed(Path(path).read_text(encoding="utf-8", errors="replace"))
    return re.sub(r"\n\s*\n+", "\n\n", "".join(p.out)).strip()


def extract(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    if ext in (".doc", ".rtf", ".odt"):
        c = _convert(path, "docx")
        return _docx(c) if c else "(Could not convert this old Word file.)"
    if ext == ".ppt":
        c = _convert(path, "pptx")
        return _pptx(c) if c else "(Could not convert this old PowerPoint file.)"
    if ext == ".xls":
        c = _convert(path, "xlsx")
        return _xlsx(c) if c else "(Could not convert this old Excel file.)"
    if ext in (".html", ".htm"):
        return _html(path)
    if ext in IMG_EXT:
        return f"(Image {path.name} - open the file to see it.)"
    if ext in AUDIO_EXT:
        return f"(Recording {path.name} - transcribed automatically overnight into Notes/Transcripts.)"
    if ext == ".zip":
        z = zipfile.ZipFile(path)
        return "ZIP file containing:\n" + "\n".join(f"  {i.filename} ({i.file_size:,} bytes)" for i in z.infolist()
                                                     if not i.is_dir()) + "\nRun school.py sync to file its contents."
    raw = path.read_bytes()
    if b"\x00" in raw[:4000]:
        return f"(Binary file {path.name} - no text to read.)"
    return raw.decode("utf-8", "replace")


def read_cmd(f, part=1):
    p = Path(f).expanduser()
    if not p.exists():
        hits = find_in_library(f)
        if len(hits) == 1:
            p = hits[0]
        elif hits:
            print("Several files match - use one of these paths:\n" + "\n".join(f"  {h}" for h in hits[:10]))
            return
        else:
            sys.exit(f"No such file: {f}")
    t = extract(p)
    parts = max(1, -(-len(t) // PART_CHARS))
    part = min(max(1, part), parts)
    print(f"FILE: {p}\nPART {part} of {parts} ({len(t):,} characters)\n")
    print(t[(part - 1) * PART_CHARS: part * PART_CHARS])
    if part < parts:
        print(f"\n[... more: run  school.py read \"{p}\" {part + 1}]")


# ---------------------------------------------------------------- library
def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def load_reg():
    try:
        return json.load(open(REG))
    except (OSError, ValueError):
        return {"files": {}}


def save_reg(r):
    LIB.mkdir(parents=True, exist_ok=True)
    json.dump(r, open(REG, "w"), indent=1)


def clean_name(name):
    n = re.sub(r"^doc_[0-9a-f]{8,}_", "", name)
    n = re.sub(r"\s*\((\d)\)(?=\.\w+$)", "", n)          # "file (1).pdf" -> "file.pdf"
    n = re.sub(r"\s+_(\d)_(?=\.\w+$)", "", n)
    n = n.replace("—", "-").replace("/", "-")
    if re.match(r"26_FA_", n):
        code = re.search(r"26_FA_([A-Z]+)_(\d+)", n)
        n = f"{code.group(1)} {code.group(2)} Syllabus (Fall 2026).pdf" if code else n
    if n.lower() == "message.txt":
        n = "Transcript.txt"
    return n


def guess_course(name, text):
    name = name.replace("_", " ")
    head = name + "\n" + text[:6000]
    for c, info in COURSES.items():
        if re.search(info["code"], name, re.I):
            return c
    for c, info in COURSES.items():
        if re.search(info["code"], head, re.I):
            return c
    scores = {c: len(re.findall(info["hints"], head, re.I)) for c, info in COURSES.items()}
    best = max(scores, key=scores.get)
    return best if scores[best] else "_Unsorted"


def guess_folder(name, text):
    low, head = name.lower(), text[:3000].lower()
    if "syllabus" in low or ("syllabus" in head and "concourse" in head):
        return "Syllabus"
    if "transcript" in low or head.startswith("here is the transcript") or low.endswith((".md",)) and "notes" in low:
        return "Notes"
    if re.search(r"instruction|reference|rubric|sample|topology|template|resources|guide|reverse eng|final project\.pdf|"
                 r"extra credit", low) or low.endswith((".csv", ".sql", ".json")):
        return "Course Materials"
    if low.endswith((".docx", ".pptx")) or re.search(r"submission|brief|portfolio|draft|assignment \d", low):
        return "My Work"
    return "Course Materials"


def text_path(dest):
    return TEXT / dest.relative_to(LIB).with_suffix(dest.suffix + ".txt")


def add_file(src, course=None, folder=None, reg=None, label=None):
    """Copy one file into the library (skips exact duplicates). Returns (status, dest)."""
    own = reg is None
    reg = reg or load_reg()
    h = sha(src)
    if h in reg["files"] and Path(reg["files"][h]["path"]).exists():
        return "duplicate", Path(reg["files"][h]["path"])
    name = clean_name(label or src.name)
    try:
        text = extract(src)
    except Exception as ex:
        text = f"(Could not read: {type(ex).__name__}: {ex})"
    course = course or guess_course(name, text)
    folder = folder or guess_folder(name, text)
    dest = LIB / course / folder / name
    n = 2
    while dest.exists():
        dest = dest.with_name(f"{Path(name).stem} ({n}){Path(name).suffix}")
        n += 1
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    tp = text_path(dest)
    tp.parent.mkdir(parents=True, exist_ok=True)
    tp.write_text(text, encoding="utf-8")
    reg["files"][h] = {"path": str(dest), "source": str(src), "added": str(dt.date.today()), "course": course}
    if own:
        save_reg(reg)
    return "added", dest


def sync():
    reg = load_reg()
    seen = {v["source"] for v in reg["files"].values()}
    added = []
    for root in SOURCES:
        if not root.is_dir():
            continue
        for p in sorted(root.rglob("*")):
            if not p.is_file() or p.name.startswith(".") or str(p) in seen:
                continue
            ext = p.suffix.lower()
            if ext == ".zip":
                with tempfile.TemporaryDirectory() as tmp:
                    try:
                        zipfile.ZipFile(p).extractall(tmp)
                    except Exception as ex:
                        print(f"Could not open zip {p.name}: {ex}")
                        continue
                    for q in sorted(Path(tmp).rglob("*")):
                        if q.is_file() and not q.name.startswith((".", "__MACOSX")) and "__MACOSX" not in str(q):
                            if not q.suffix:                    # no extension: keep only if it's really a PDF
                                label = q.name + ".pdf" if q.read_bytes()[:4] == b"%PDF" else None
                                if not label:
                                    continue
                            elif q.suffix.lower() not in DOC_EXT:  # icons/images inside saved web pages: skip
                                continue
                            else:
                                label = None
                            st, d = add_file(q, reg=reg, label=label)
                            reg["files"][sha(q)]["source"] = f"{p}!{q.relative_to(tmp)}"
                            if st == "added":
                                added.append(d)
                reg["files"].setdefault("zip:" + str(p), {"path": "", "source": str(p), "added": str(dt.date.today())})
                continue
            if ext not in DOC_EXT | IMG_EXT | AUDIO_EXT:
                continue
            if ext in IMG_EXT | AUDIO_EXT:                 # course can't be guessed from a photo: date-stamp it in _Unsorted
                kind = "Photos" if ext in IMG_EXT else "Recordings"
                stamp = dt.datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H%M")
                st, d = add_file(p, course="_Unsorted", folder=kind, reg=reg, label=f"{stamp} {p.name}")
            else:
                st, d = add_file(p, reg=reg)
            if st == "added":
                added.append(d)
            else:
                reg["files"]["dup:" + str(p)] = {"path": str(d), "source": str(p), "added": str(dt.date.today())}
    save_reg(reg)
    index()
    if not added:
        print("Library is up to date - no new files.")
        return
    print(f"Filed {len(added)} new file(s):")
    for d in added:
        print(f"  {d.relative_to(LIB)}")
    if any("_Unsorted" in str(d) for d in added):
        print('Some files went to _Unsorted (photos and voice notes always do). Look at them, then file each one: '
              'school.py file "<path>" <COURSE> [folder]')


def file_cmd(f, course, folder=None):
    course = next((c for c in list(COURSES) + ["_Unsorted"] if c.lower().replace("-", "") == course.lower().replace("-", "").replace(" ", "")), None)
    if not course:
        sys.exit(f"Unknown course. Use one of: {', '.join(COURSES)}")
    folder = next((x for x in FOLDERS if folder and x.lower().startswith(folder.lower()[:4])), folder)
    p = Path(f).expanduser()
    if not p.exists():
        hits = find_in_library(f)
        if len(hits) != 1:
            sys.exit(f"Found {len(hits)} files matching '{f}'. Give the full path.")
        p = hits[0]
    reg = load_reg()
    if LIB in p.parents:  # move an already-filed file
        folder = folder or p.parent.name
        dest = LIB / course / folder / p.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(p), str(dest))
        old_t, new_t = text_path(p), text_path(dest)
        if old_t.exists():
            new_t.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(old_t), str(new_t))
        for v in reg["files"].values():
            if v.get("path") == str(p):
                v["path"], v["course"] = str(dest), course
        save_reg(reg)
        print(f"Moved to {dest.relative_to(LIB)}")
    else:
        st, d = add_file(p, course, folder, reg)
        save_reg(reg)
        print(("Already in the library: " if st == "duplicate" else "Filed: ") + str(d.relative_to(LIB)))
    index()


def find_in_library(q):
    q = q.lower()
    return [p for p in LIB.rglob("*") if p.is_file() and ".text" not in p.parts and not p.name.startswith(".")
            and q in p.name.lower()]


# ---------------------------------------------------------------- search
STOP = set("the and for with that this from what when where which about into have has are was were will your you how does "
           "did can could should would its their them they then than there these those been being also just more most".split())


def search(query, course=None):
    words = [w for w in re.findall(r"[a-z0-9+#.]{3,}", query.lower()) if w not in STOP]
    if not words:
        sys.exit("Give me a few key words to search for.")
    phrase = query.lower().strip()
    results = []
    roots = [TEXT / c for c in ([course] if course else os.listdir(TEXT) if TEXT.exists() else [])]
    notes = [LIB / c / "Notes" for c in ([course] if course else COURSES)]
    files = [f for r in roots if Path(r).is_dir() for f in Path(r).rglob("*.txt")]
    files += [f for r in notes if r.is_dir() for f in r.rglob("*.md")]
    for f in files:
        t = f.read_text(errors="replace")
        page_marks = [(m.start(), m.group(1)) for m in re.finditer(r"=== (?:page|slide|sheet) ([^=]+?) ===", t)]
        for m in re.finditer(r"(?s).{1,700}?(?:\n\n|\n(?==== )|$)", t):
            chunk = m.group(0)
            low = chunk.lower()
            hits = sum(1 for w in set(words) if w in low)
            if not hits:
                continue
            score = hits * 10 + (25 if phrase in low else 0) + sum(low.count(w) for w in words)
            page = next((pg for pos, pg in reversed(page_marks) if pos <= m.start()), None)
            results.append((score, f, page, chunk))
    if not results:
        print(f"Nothing in the library mentions: {', '.join(words)}. Try other words, or check Documents/Pi files.")
        return
    results.sort(key=lambda r: -r[0])
    print(f"Top matches for: {query}\n")
    shown = 0
    for score, f, page, chunk in results:
        src = f.relative_to(TEXT).with_suffix("") if TEXT in f.parents else f.relative_to(LIB)
        snippet = re.sub(r"\s+", " ", chunk).strip()
        if len(snippet) > 420:
            i = min((snippet.lower().find(w) for w in words if w in snippet.lower()), default=0)
            snippet = ("..." if i > 150 else "") + snippet[max(0, i - 150): i + 270] + "..."
        print(f"[{shown + 1}] {src}" + (f" (page {page.strip()})" if page else ""))
        print(f"    {snippet}\n")
        shown += 1
        if shown == 6:
            break
    print(f"Full file: school.py read \"{LIB}/<path above>\"")


# ---------------------------------------------------------------- index
def index():
    LIB.mkdir(parents=True, exist_ok=True)
    now = dt.datetime.now()
    top = ["# Athena - School Library", "", f"_Rebuilt {now:%Y-%m-%d %H:%M} by school.py. Don't edit by hand._", "",
           "Every school file lives here, one folder per course. Search everything with "
           "`school.py search \"words\"`.", "", "| Course | Files | Last added |", "|---|---|---|"]
    for c in list(COURSES) + ["_Unsorted"]:
        d = LIB / c
        files = [p for p in d.rglob("*") if p.is_file() and not p.name.startswith((".", "00 Index"))] if d.is_dir() else []
        if not files and c == "_Unsorted":
            continue
        last = max((p.stat().st_mtime for p in files), default=0)
        top.append(f"| [[{c}/00 Index|{c}]] {COURSES.get(c, {}).get('name', 'not sorted yet')} | {len(files)} | "
                   f"{dt.datetime.fromtimestamp(last):%Y-%m-%d} |" if last else f"| {c} | 0 | - |")
        if c in COURSES or files:
            d.mkdir(parents=True, exist_ok=True)
            L = [f"# {c} - {COURSES.get(c, {}).get('name', 'Unsorted')}", "",
                 f"_Rebuilt {now:%Y-%m-%d %H:%M} by school.py._", ""]
            for fo in FOLDERS if c in COURSES else sorted({p.parent.name for p in files}):
                fl = sorted((p for p in files if p.parent.name == fo), key=lambda p: p.name.lower())
                L.append(f"## {fo}")
                L += [f"- {'[[' + p.stem + ']]' if p.suffix == '.md' else '`' + p.name + '`'}"
                      f"  ({p.stat().st_size / 1024:.0f} KB, {dt.datetime.fromtimestamp(p.stat().st_mtime):%b %d})"
                      for p in fl] or ["- (nothing yet)"]
                L.append("")
            (d / "00 Index.md").write_text("\n".join(L), encoding="utf-8")
    cal = PROFILE / "school/deadlines.json"
    if cal.exists():
        top += ["", "Deadlines: see `school/deadlines.json` via `planner.py list` (Athena's planner)."]
    (LIB / "00 Index.md").write_text("\n".join(top) + "\n", encoding="utf-8")


def list_cmd(course=None):
    for c in ([course] if course else list(COURSES) + ["_Unsorted"]):
        d = LIB / c
        if not d.is_dir():
            continue
        files = sorted(p for p in d.rglob("*") if p.is_file() and not p.name.startswith((".", "00 Index")))
        print(f"{c} ({len(files)} files)")
        for p in files:
            print(f"  {p.relative_to(d)}")


if __name__ == "__main__":
    a = sys.argv[1:]
    try:
        if not a:
            print(__doc__)
        elif a[0] == "sync":
            sync()
        elif a[0] == "read" and len(a) > 1:
            read_cmd(a[1], int(a[2]) if len(a) > 2 and a[2].isdigit() else 1)
        elif a[0] == "search" and len(a) > 1:
            c = a[2] if len(a) > 2 and a[2].upper().replace(" ", "-") in COURSES else None
            q = " ".join(a[1:2] + (a[2:] if not c else a[3:]))
            search(q, c.upper().replace(" ", "-") if c else None)
        elif a[0] == "list":
            list_cmd(a[1].upper() if len(a) > 1 else None)
        elif a[0] == "file" and len(a) > 2:
            file_cmd(a[1], a[2], a[3] if len(a) > 3 else None)
        elif a[0] == "index":
            index()
            print(f"Index rebuilt: {LIB / '00 Index.md'}")
        else:
            print(__doc__)
    except SystemExit:
        raise
    except Exception as ex:
        print(f"SCHOOL ERROR: {type(ex).__name__}: {ex}")
        sys.exit(1)
